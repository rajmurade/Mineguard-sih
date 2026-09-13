"""PDF (ReportLab) and XLSX (OpenPyXL) report generation for MineGuard.

Pulls incidents + environmental readings within a ``daily`` (24h) or
``weekly`` (7d) range and summarizes them by severity, type, status and zone.
"""

from __future__ import annotations

from collections import Counter
from datetime import datetime, timedelta, timezone
from io import BytesIO

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.enums import ResolutionStatus, Severity
from app.models import EnvironmentalReading, Incident

RANGE_HOURS = {"daily": 24, "weekly": 24 * 7}

FORMAT_MEDIA = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

SEVERITY_ORDER = [Severity.compliant, Severity.warning, Severity.high, Severity.critical]


def _hours_for(range_: str) -> int:
    try:
        return RANGE_HOURS[range_]
    except KeyError:
        raise ValueError(f"unknown report range: {range_} (expected daily|weekly)")


class ReportData:
    """Timed + summarized view of the DB for one report."""

    def __init__(self, db: Session, range_: str):
        self.range_ = range_
        self.now = datetime.now(timezone.utc)
        self.cutoff = self.now - timedelta(hours=_hours_for(range_))

        self.incidents = db.scalars(
            select(Incident)
            .where(Incident.timestamp >= self.cutoff)
            .order_by(Incident.timestamp.desc())
        ).all()
        self.readings = db.scalars(
            select(EnvironmentalReading)
            .where(EnvironmentalReading.timestamp >= self.cutoff)
            .order_by(EnvironmentalReading.timestamp.desc())
        ).all()

    @property
    def incidents_by_severity(self) -> list[tuple[Severity, int]]:
        counts = Counter(i.severity for i in self.incidents)
        return [(sev, counts.get(sev, 0)) for sev in SEVERITY_ORDER]

    @property
    def incidents_by_type(self) -> list[tuple[str, int]]:
        return list(Counter(i.type.value for i in self.incidents).items())

    @property
    def incidents_by_status(self) -> list[tuple[str, int]]:
        return list(
            Counter(i.resolution_status.value for i in self.incidents).items()
        )

    @property
    def latest_reading_per_zone(self) -> list[EnvironmentalReading]:
        latest: dict[str, EnvironmentalReading] = {}
        for reading in self.readings:  # already newest-first
            latest.setdefault(reading.zone, reading)
        return list(latest.values())


# ----------------------------------------------------------------------
# PDF export (ReportLab platypus)
# ----------------------------------------------------------------------


def _pdf(data: ReportData) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title="MineGuard Report",
    )
    styles = getSampleStyleSheet()

    def table_from_rows(headers: list[str], rows: list[tuple]) -> Table:
        body = [headers, *[list(r) for r in rows]]
        table = Table(body, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94a3b8")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        return table

    elements = [
        Paragraph(f"MineGuard — {data.range_.title()} Report", styles["Title"]),
        Spacer(1, 4 * mm),
        Paragraph(
            f"Generated {data.now.strftime('%Y-%m-%d %H:%M UTC')} · window: "
            f"{data.cutoff.strftime('%Y-%m-%d %H:%M UTC')} → now "
            f"({len(data.incidents)} incidents, {len(data.readings)} environmental readings)",
            styles["Normal"],
        ),
        Spacer(1, 6 * mm),
        Paragraph("Incidents by severity", styles["Heading2"]),
        table_from_rows(
            ["Severity", "Count"],
            [(sev.value, n) for sev, n in data.incidents_by_severity],
        ),
        Spacer(1, 4 * mm),
        Paragraph("Incidents by type", styles["Heading2"]),
        table_from_rows(["Type", "Count"], data.incidents_by_type),
        Spacer(1, 4 * mm),
        Paragraph("Incidents by resolution status", styles["Heading2"]),
        table_from_rows(["Status", "Count"], data.incidents_by_status),
        Spacer(1, 4 * mm),
        Paragraph("Latest environmental readings per zone", styles["Heading2"]),
        table_from_rows(
            ["Zone", "AQI", "Gas (ppm)", "Temp (°C)", "Noise (dB)", "Dust (mg/m³)"],
            [
                (
                    r.zone,
                    r.aqi,
                    r.gas_level,
                    r.temperature,
                    r.noise_db,
                    r.dust_pm,
                )
                for r in data.latest_reading_per_zone
            ],
        ),
        Spacer(1, 6 * mm),
        Paragraph("Incident log", styles["Heading2"]),
        table_from_rows(
            ["ID", "Time (UTC)", "Type", "Severity", "Zone", "Status", "Description"],
            [
                (
                    i.id,
                    i.timestamp.strftime("%Y-%m-%d %H:%M") if i.timestamp else "",
                    i.type.value,
                    i.severity.value,
                    i.zone,
                    i.resolution_status.value,
                    (i.description or "")[:90],
                )
                for i in data.incidents
            ],
        ),
    ]

    doc.build(elements)
    return buffer.getvalue()


# ----------------------------------------------------------------------
# XLSX export (OpenPyXL)
# ----------------------------------------------------------------------


def _xlsx(data: ReportData) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()

    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = f"MineGuard {data.range_.title()} Report"
    summary["A2"] = f"Generated {data.now.isoformat()}"
    summary["A3"] = f"Window start: {data.cutoff.isoformat()}"
    summary["A5"] = "Incidents by severity"
    summary.append(["Severity", "Count"])
    for sev, count in data.incidents_by_severity:
        summary.append([sev.value, count])
    row = summary.max_row + 2
    summary.cell(row=row, column=1, value="Incidents by type").font = Font(bold=True)
    summary.append(["Type", "Count"])
    for event_type, count in data.incidents_by_type:
        summary.append([event_type, count])
    row = summary.max_row + 2
    summary.cell(row=row, column=1, value="Incidents by status").font = Font(bold=True)
    summary.append(["Status", "Count"])
    for status, count in data.incidents_by_status:
        summary.append([status, count])

    incidents_sheet = workbook.create_sheet("Incidents")
    incidents_sheet.append(["ID", "Timestamp (UTC)", "Type", "Severity", "Zone", "Status", "Worker ID", "Description"])
    for i in data.incidents:
        incidents_sheet.append(
            [
                i.id,
                i.timestamp.replace(tzinfo=None) if i.timestamp else None,
                i.type.value,
                i.severity.value,
                i.zone,
                i.resolution_status.value,
                i.worker_id,
                i.description,
            ]
        )

    readings_sheet = workbook.create_sheet("Environmental")
    readings_sheet.append(
        ["Zone", "Timestamp (UTC)", "AQI", "Gas (ppm)", "Temp (°C)", "Humidity (%)", "Noise (dB)", "Dust (mg/m³)"]
    )
    for r in data.readings:
        readings_sheet.append(
            [
                r.zone,
                r.timestamp.replace(tzinfo=None) if r.timestamp else None,
                r.aqi,
                r.gas_level,
                r.temperature,
                r.humidity,
                r.noise_db,
                r.dust_pm,
            ]
        )

    header_fill = PatternFill("solid", fgColor="0F172A")
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ----------------------------------------------------------------------
# entry point
# ----------------------------------------------------------------------


def export_report(db: Session, format_: str, range_: str) -> tuple[bytes, str, str]:
    """Generate a report. Returns ``(bytes, filename, media_type)``."""
    if format_ not in FORMAT_MEDIA:
        raise ValueError(f"unknown report format: {format_} (expected pdf|xlsx)")
    data = ReportData(db, range_)
    payload = _pdf(data) if format_ == "pdf" else _xlsx(data)
    stamp = data.now.strftime("%Y%m%d-%H%M%S")
    filename = f"mineguard_{range_}_{stamp}.{format_}"
    return payload, filename, FORMAT_MEDIA[format_]