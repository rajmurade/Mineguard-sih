"""Smoke tests for /reports/export and the incident evidence file endpoint.

Reuses the SQLite TestClient harness from test_cv_backend so both modules
share the same engine + get_db override (a second module-level override would
shadow the first and drop its tables).
"""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from io import BytesIO

from sqlalchemy import select

from app.enums import IncidentType, ResolutionStatus, Severity
from app.models import EnvironmentalReading, Incident
from test_cv_backend import TestingSessionLocal, client  # noqa: F401 (shared harness)


def _insert_incident(db, evidence_path=None):
    incident = Incident(
        type=IncidentType.ppe_violation,
        zone="Zone A - Tunnel 1",
        timestamp=datetime.now(timezone.utc) - timedelta(hours=2),
        severity=Severity.high,
        worker_id=None,
        evidence_frame_path=evidence_path,
        resolution_status=ResolutionStatus.open,
        description="smoke test incident",
    )
    db.add(incident)
    db.commit()
    db.refresh(incident)
    return incident


def test_export_pdf():
    with TestingSessionLocal() as db:
        _insert_incident(db)
    response = client.get("/reports/export?format=pdf&range=daily")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/pdf")
    assert "Content-Disposition" in response.headers
    assert len(response.content) > 1000  # real PDF with the seeded incident


def test_export_xlsx_is_valid_spreadsheet():
    with TestingSessionLocal() as db:
        _insert_incident(db)
    response = client.get("/reports/export?format=xlsx&range=weekly")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats"
    )
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Summary", "Incidents", "Environmental"]
    # seeded incident appears in the Incidents sheet
    rows = list(workbook["Incidents"].iter_rows(values_only=True))
    assert any("smoke test incident" in str(r) for r in rows)


def test_export_invalid_format_or_range():
    assert client.get("/reports/export?format=docx&range=daily").status_code == 422
    assert client.get("/reports/export?format=pdf&range=yearly").status_code == 422


def test_evidence_endpoint_serves_file(tmp_path):
    evidence = tmp_path / "frame.jpg"
    evidence.write_bytes(b"\xff\xd8\xff\xe0fake-jpeg-content")
    with TestingSessionLocal() as db:
        incident = _insert_incident(db, evidence_path=str(evidence))

    response = client.get(f"/incidents/{incident.id}/evidence")
    assert response.status_code == 200
    assert response.content == b"\xff\xd8\xff\xe0fake-jpeg-content"
    assert response.headers["content-type"].startswith("image/")


def test_evidence_endpoint_resolves_cwd_relative_path(tmp_path, monkeypatch):
    """A stored cwd-relative path (evidence/no-vest/x.jpg) resolves against
    the working directory, mirroring the API container's ./evidence:/app/evidence
    mount where the CV pipeline writes violation frames."""
    evidence_file = tmp_path / "evidence" / "no-vest" / "frame.jpg"
    evidence_file.parent.mkdir(parents=True)
    evidence_file.write_bytes(b"\xff\xd8\xff\xe0fake-jpeg-content")
    monkeypatch.chdir(tmp_path)

    with TestingSessionLocal() as db:
        incident = _insert_incident(db, evidence_path="evidence/no-vest/frame.jpg")

    response = client.get(f"/incidents/{incident.id}/evidence")
    assert response.status_code == 200
    assert response.content == b"\xff\xd8\xff\xe0fake-jpeg-content"


def test_evidence_endpoint_404_without_frame():
    with TestingSessionLocal() as db:
        incident = _insert_incident(db, evidence_path=None)
    assert client.get(f"/incidents/{incident.id}/evidence").status_code == 404


def test_evidence_endpoint_404_for_missing_file(tmp_path):
    with TestingSessionLocal() as db:
        incident = _insert_incident(
            db, evidence_path=str(tmp_path / "not-there.jpg")
        )
    assert client.get(f"/incidents/{incident.id}/evidence").status_code == 404